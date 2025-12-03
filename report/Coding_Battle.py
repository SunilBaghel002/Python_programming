from docx import Document
from docx.shared import Inches

doc = Document()

# Title
doc.add_heading("CodeBattle Event Report – Aravali College of Engineering & Management", level=1)

# Circular
doc.add_heading("CIRCULAR", level=2)
doc.add_paragraph(
    "Subject: Coding Battle – “GET SET CODE” on 30th July 2025\n\n"
    "All B.Tech students of Aravali College of Engineering & Management are informed that the "
    "department is organizing a coding competition titled *TECHNITUDE’s Coding Battle* on "
    "30th July 2025 at 2:20 PM in Lab 9.\n\n"
    "All interested students are encouraged to participate and showcase their programming skills.\n\n"
    "Sincerely,\nRegistrar\nAravali College of Engineering and Management"
)

# Event Brochure
doc.add_heading("EVENT BROCHURE", level=2)
doc.add_paragraph(
    "Event: TECHNITUDE’s Coding Battle – “GET SET CODE”\n"
    "Date: 30th July 2025\n"
    "Time: 2:20 PM\n"
    "Venue: Lab 9, ACEM\n\n"
    "Event Schedule:\n"
    "• 02:00 PM – Reporting & Registration\n"
    "• 02:20 PM – Round 1 (Debugging Round)\n"
    "• 03:00 PM – Round 2 (Problem-Solving Round)\n"
    "• 04:00 PM – Final Round\n"
    "• 04:40 PM – Result Announcement & Prize Distribution\n\n"
    "Technical Head: Mr. Umesh Goyal\n"
    "Head Coordinator: Satyam Pandey (8851020767)"
)

# Event Poster
doc.add_heading("Event Poster", level=2)
try:
    doc.add_picture('/mnt/data/WhatsApp Image 2025-07-29 at 15.25.15_311c710c.jpg', width=Inches(5))
except:
    doc.add_paragraph("[Poster Image Attached Separately]")

# Event Report
doc.add_heading("EVENT REPORT", level=2)
doc.add_paragraph(
    "The *TECHNITUDE’s Coding Battle – GET SET CODE* was successfully conducted on "
    "30th July 2025 in Lab 9 at 2:20 PM. The event witnessed active participation from B.Tech "
    "students of ACEM. The competition aimed to enhance competitive programming skills, logical "
    "thinking, and real-time debugging abilities.\n\n"
    "The event comprised three rounds: Debugging, Problem Solving, and a Final Coding Challenge. "
    "Participants demonstrated exceptional enthusiasm and teamwork throughout the competition. "
    "Judges evaluated participants on logic, accuracy, execution time, and code quality.\n\n"
    "The event concluded with a prize distribution ceremony where the top performers were "
    "awarded certificates and recognition.\n\n"
)

# Participants List
doc.add_heading("LIST OF PARTICIPANTS", level=2)
table = doc.add_table(rows=1, cols=3)
hdr = table.rows[0].cells
hdr[0].text = "S.No."
hdr[1].text = "Name of Student"
hdr[2].text = "College Name"

for i in range(1, 31):
    row = table.add_row().cells
    row[0].text = str(i)
    row[1].text = ""
    row[2].text = "ACEM"

# Feedback Form
doc.add_heading("FEEDBACK FORM", level=2)
doc.add_paragraph(
    "Name: __________________________\n"
    "Designation (Student/Faculty/Staff): __________________________\n"
    "Email: __________________________\n\n"
    "How satisfied were you with the event?\n"
    "☐ Very Satisfied   ☐ Satisfied   ☐ Neutral   ☐ Dissatisfied   ☐ Very Dissatisfied\n\n"
    "How would you rate the fairness of evaluation?\n"
    "☐ Excellent   ☐ Good   ☐ Average   ☐ Poor\n\n"
    "Would you like to participate in similar events again?\n"
    "☐ Yes   ☐ No\n\n"
    "What did you like the most about the event?\n"
    "_____________________________________________________________\n\n"
    "Suggestions for improvement:\n"
    "_____________________________________________________________\n\n"
    "Google Form Link:\n"
    "https://docs.google.com/forms/d/e/1FAIpQLScwgzIzaiBOT47A2_-FJ7FDRGIJIUnfCrqrss9QTlHslJZtfg/viewform"
)

# Save the file
path = "/mnt/data/CodeBattle_Event_Report.docx"
doc.save(path)

path
